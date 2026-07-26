#!/usr/bin/env python3
"""Reconcile signed V2 movements with the adjusted PDET monthly series."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import tempfile
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
import requests
from openpyxl import load_workbook


PACKAGE_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = (
    PACKAGE_ROOT
    / "data"
    / "vintage"
    / "pdet"
    / "3-tabelas_Maio_de_2026.xlsx"
)
DEFAULT_MANIFEST = PACKAGE_ROOT / "data" / "vintage" / "manifest.json"
DEFAULT_MOVEMENTS = (
    PACKAGE_ROOT
    / "data"
    / "interim"
    / "movimentacoes"
    / "competenciamov=*"
    / "part.parquet"
)
DEFAULT_OUTPUT = (
    PACKAGE_ROOT
    / "results"
    / "reconciliation"
    / "pdet_vs_v2_mensal.csv"
)
DEFAULT_SUPPORT = (
    PACKAGE_ROOT
    / "results"
    / "reconciliation"
    / "pdet_vs_v2_support.json"
)
PDET_SOURCE_PAGE = (
    "https://www.gov.br/trabalho-e-emprego/pt-br/acesso-a-informacao/"
    "acoes-e-programas/programas-projetos-acoes-obras-e-atividades/"
    "estatisticas-trabalho/novo-caged/2026/maio/pagina-inicial"
)
PDET_DRIVE_FOLDER = (
    "https://drive.google.com/drive/folders/"
    "1F89h6odTPGIGMb9eDiJKCute9W89QmqN?usp=sharing"
)
PDET_FILE_ID = "12DjeOFjIgeJNp-3OxslJqJNNmA1ViQMj"
PDET_DOWNLOAD_URL = (
    "https://drive.usercontent.google.com/download"
    f"?id={PDET_FILE_ID}&export=download&confirm=t"
)
START_PERIOD = 202101
END_PERIOD = 202605
MONTHS = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def utc_now() -> str:
    return (
        datetime.now(timezone.utc)
        .isoformat()
        .replace("+00:00", "Z")
    )


def _atomic_json(payload: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(f"{path.suffix}.tmp")
    temporary.write_text(
        json.dumps(
            payload,
            indent=2,
            sort_keys=True,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    os.replace(temporary, path)


def _read_manifest(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Vintage manifest must be a JSON object")
    return payload


def _copy_atomic(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(f"{destination.suffix}.tmp")
    shutil.copyfile(source, temporary)
    os.replace(temporary, destination)


def register_pdet_workbook(
    source: Path,
    destination: Path = DEFAULT_WORKBOOK,
    manifest_path: Path = DEFAULT_MANIFEST,
    *,
    accessed_at: str | None = None,
) -> dict[str, Any]:
    if not source.is_file():
        raise FileNotFoundError(f"PDET workbook not found: {source}")
    if source.resolve() != destination.resolve():
        _copy_atomic(source, destination)
    with destination.open("rb") as handle:
        signature = handle.read(4)
    if signature != b"PK\x03\x04":
        raise RuntimeError("PDET source is not an XLSX workbook")
    parse_adjusted_series(destination)
    entry = {
        "accessed_at": accessed_at or utc_now(),
        "bytes": destination.stat().st_size,
        "drive_file_id": PDET_FILE_ID,
        "drive_folder": PDET_DRIVE_FOLDER,
        "series": "Tabela 5.1 - com ajustes",
        "sha256": sha256_file(destination),
        "source_page": PDET_SOURCE_PAGE,
        "url": PDET_DOWNLOAD_URL,
    }
    manifest = _read_manifest(manifest_path)
    manifest[f"pdet/{destination.name}"] = entry
    _atomic_json(manifest, manifest_path)
    return entry


def download_pdet_workbook(
    destination: Path = DEFAULT_WORKBOOK,
    manifest_path: Path = DEFAULT_MANIFEST,
) -> dict[str, Any]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix="pdet-",
        suffix=".xlsx",
        dir=destination.parent,
        delete=False,
    ) as handle:
        temporary = Path(handle.name)
    try:
        with requests.get(
            PDET_DOWNLOAD_URL,
            stream=True,
            timeout=120,
        ) as response:
            response.raise_for_status()
            with temporary.open("wb") as handle:
                for chunk in response.iter_content(1024 * 1024):
                    if chunk:
                        handle.write(chunk)
        return register_pdet_workbook(
            temporary,
            destination,
            manifest_path,
        )
    finally:
        temporary.unlink(missing_ok=True)


def _normalize_text(value: Any) -> str:
    return "".join(
        character
        for character in unicodedata.normalize(
            "NFKD",
            str(value).strip().lower(),
        )
        if not unicodedata.combining(character)
    )


def _period_from_label(value: Any) -> int | None:
    normalized = _normalize_text(value)
    if "/" not in normalized:
        return None
    month_name, year_text = normalized.split("/", maxsplit=1)
    if month_name not in MONTHS or not year_text.isdigit():
        return None
    return int(year_text) * 100 + MONTHS[month_name]


def _expected_periods(start_period: int, end_period: int) -> list[int]:
    start = pd.Period(str(start_period), freq="M")
    end = pd.Period(str(end_period), freq="M")
    return [
        int(period.strftime("%Y%m"))
        for period in pd.period_range(start, end, freq="M")
    ]


def parse_adjusted_series(
    workbook_path: Path,
    *,
    start_period: int = START_PERIOD,
    end_period: int = END_PERIOD,
) -> pd.DataFrame:
    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )
    try:
        if "Tabela 5.1" not in workbook.sheetnames:
            raise RuntimeError(
                "PDET workbook has no adjusted-series sheet Tabela 5.1"
            )
        sheet = workbook["Tabela 5.1"]
        records: list[dict[str, int]] = []
        for row in sheet.iter_rows(min_row=6, values_only=True):
            period = _period_from_label(row[1] if len(row) > 1 else None)
            if period is None or not start_period <= period <= end_period:
                continue
            admissions = int(row[3])
            separations = int(row[4])
            balance = int(row[5])
            records.append(
                {
                    "competenciamov": period,
                    "pdet_admissoes": admissions,
                    "pdet_desligamentos": separations,
                    "pdet_saldo": balance,
                }
            )
    finally:
        workbook.close()

    series = pd.DataFrame(records).sort_values(
        "competenciamov"
    ).reset_index(drop=True)
    expected = _expected_periods(start_period, end_period)
    if series["competenciamov"].tolist() != expected:
        raise RuntimeError(
            "Adjusted PDET series does not cover every expected month"
        )
    if (
        series["pdet_admissoes"] - series["pdet_desligamentos"]
        != series["pdet_saldo"]
    ).any():
        raise RuntimeError("Adjusted PDET series violates balance identity")
    return series


def aggregate_v2(
    movements_glob: Path,
    *,
    start_period: int = START_PERIOD,
    end_period: int = END_PERIOD,
) -> pd.DataFrame:
    source = "'" + str(movements_glob).replace("'", "''") + "'"
    connection = duckdb.connect()
    try:
        invalid = int(
            connection.execute(
                f"""
                SELECT count(*)
                FROM read_parquet({source}, hive_partitioning = true)
                WHERE CAST(competenciamov AS INTEGER)
                      BETWEEN {start_period} AND {end_period}
                  AND saldomovimentacao NOT IN (-1, 1)
                """
            ).fetchone()[0]
        )
        if invalid:
            raise RuntimeError(
                f"V2 contains {invalid} invalid movement signs"
            )
        frame = connection.execute(
            f"""
            SELECT
                CAST(competenciamov AS INTEGER) AS competenciamov,
                CAST(sum(
                    CASE WHEN saldomovimentacao = 1
                         THEN peso ELSE 0 END
                ) AS BIGINT) AS v2_admissoes,
                CAST(sum(
                    CASE WHEN saldomovimentacao = -1
                         THEN peso ELSE 0 END
                ) AS BIGINT) AS v2_desligamentos
            FROM read_parquet({source}, hive_partitioning = true)
            WHERE CAST(competenciamov AS INTEGER)
                  BETWEEN {start_period} AND {end_period}
            GROUP BY competenciamov
            ORDER BY competenciamov
            """
        ).df()
    finally:
        connection.close()
    frame["v2_saldo"] = (
        frame["v2_admissoes"] - frame["v2_desligamentos"]
    )
    expected = _expected_periods(start_period, end_period)
    if frame["competenciamov"].tolist() != expected:
        raise RuntimeError("V2 does not cover every expected fact month")
    return frame


def _shift_periods(values: pd.Series, shift: int) -> pd.Series:
    periods = pd.PeriodIndex(
        values.astype(str),
        freq="M",
    ) + shift
    return pd.Series(
        [int(period.strftime("%Y%m")) for period in periods],
        index=values.index,
    )


def _shift_score(
    official: pd.DataFrame,
    v2: pd.DataFrame,
    shift: int,
) -> int:
    shifted = official.copy()
    shifted["competenciamov"] = _shift_periods(
        shifted["competenciamov"],
        shift,
    )
    joined = shifted.merge(v2, on="competenciamov", how="inner")
    return int(
        sum(
            (
                joined[f"v2_{outcome}"]
                - joined[f"pdet_{outcome}"]
            )
            .abs()
            .sum()
            for outcome in ("admissoes", "desligamentos", "saldo")
        )
    )


def reconcile_monthly(
    official: pd.DataFrame,
    v2: pd.DataFrame,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    comparison = official.merge(
        v2,
        on="competenciamov",
        how="outer",
        validate="one_to_one",
        indicator=True,
    )
    if not comparison["_merge"].eq("both").all():
        raise RuntimeError("PDET and V2 monthly coverage differ")
    comparison = comparison.drop(columns="_merge").sort_values(
        "competenciamov"
    )
    for outcome in ("admissoes", "desligamentos", "saldo"):
        signed = (
            comparison[f"v2_{outcome}"]
            - comparison[f"pdet_{outcome}"]
        )
        comparison[f"diferenca_{outcome}"] = signed
        comparison[f"diferenca_{outcome}_abs"] = signed.abs()
        official_values = comparison[f"pdet_{outcome}"]
        comparison[f"diferenca_{outcome}_pct"] = (
            signed * 100.0 / official_values.where(official_values.ne(0))
        )

    shift_scores = {
        str(shift): _shift_score(official, v2, shift)
        for shift in (-1, 0, 1)
    }
    zero_score = shift_scores["0"]
    neighbor_score = min(shift_scores["-1"], shift_scores["1"])
    systematic_shift = (
        zero_score > 0
        and neighbor_score < zero_score * 0.25
    )
    difference_columns = [
        f"diferenca_{outcome}_abs"
        for outcome in ("admissoes", "desligamentos", "saldo")
    ]
    any_difference = comparison[difference_columns].gt(0).any(axis=1)
    metrics: dict[str, Any] = {
        "months": int(len(comparison)),
        "months_with_any_difference": int(any_difference.sum()),
        "max_admission_absolute_difference": int(
            comparison["diferenca_admissoes_abs"].max()
        ),
        "max_separation_absolute_difference": int(
            comparison["diferenca_desligamentos_abs"].max()
        ),
        "max_balance_absolute_difference": int(
            comparison["diferenca_saldo_abs"].max()
        ),
        "total_admission_absolute_difference": int(
            comparison["diferenca_admissoes_abs"].sum()
        ),
        "total_separation_absolute_difference": int(
            comparison["diferenca_desligamentos_abs"].sum()
        ),
        "total_balance_absolute_difference": int(
            comparison["diferenca_saldo_abs"].sum()
        ),
        "shift_scores_total_absolute_difference": shift_scores,
        "systematic_month_shift_detected": systematic_shift,
    }
    return comparison.reset_index(drop=True), metrics


def write_artifacts(
    comparison: pd.DataFrame,
    metrics: dict[str, Any],
    output_path: Path = DEFAULT_OUTPUT,
    support_path: Path = DEFAULT_SUPPORT,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(f"{output_path.suffix}.tmp")
    comparison.to_csv(temporary, index=False)
    os.replace(temporary, output_path)
    _atomic_json(metrics, support_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Reconcile V2 with the adjusted PDET monthly series."
    )
    parser.add_argument(
        "command",
        choices=("freeze", "reconcile", "all"),
    )
    parser.add_argument("--source-workbook", type=Path)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--movements-glob",
        type=Path,
        default=DEFAULT_MOVEMENTS,
    )
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--support", type=Path, default=DEFAULT_SUPPORT)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.command in {"freeze", "all"}:
        if args.source_workbook:
            entry = register_pdet_workbook(
                args.source_workbook,
                args.workbook,
                args.manifest,
            )
        else:
            entry = download_pdet_workbook(
                args.workbook,
                args.manifest,
            )
        print(json.dumps(entry, sort_keys=True))
    if args.command in {"reconcile", "all"}:
        official = parse_adjusted_series(args.workbook)
        v2 = aggregate_v2(args.movements_glob)
        comparison, metrics = reconcile_monthly(official, v2)
        write_artifacts(
            comparison,
            metrics,
            args.output,
            args.support,
        )
        if metrics["systematic_month_shift_detected"]:
            raise RuntimeError(
                "Systematic one-month reassignment shift detected"
            )
        print(json.dumps(metrics, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
